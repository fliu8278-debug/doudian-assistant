from __future__ import annotations

import html
import re
from collections import Counter
from pathlib import Path

import openpyxl


SHOP_FILE = Path(r"C:\Users\ASUS\Desktop\productExport-98c5b2f1-7500-9b0e-1267-cca088b3b4f3.xlsx")
BRAND_FILE = Path(r"C:\Users\ASUS\Desktop\23.xlsx")
OUTPUT_FILE = Path(r"E:\zizhu\outputs\shop-sku-preview.html")


def normalize_shop_sku(sku: str, raw_category: str) -> tuple[str, str, str]:
    sku = sku.strip()
    category = decode_category(raw_category)

    if not sku:
        return "", "空 SKU", category

    if category == "鞋靴":
        match = re.match(r"^(.+-[A-Za-z]+)\d+$", sku)
        return (match.group(1) if match else sku), "鞋靴：去掉末尾尺码数字", category

    if category in {"服装", "箱包"}:
        left, sep, right = sku.partition("-")
        return (f"{left}-{right[:4]}" if sep and len(right) >= 4 else sku), f"{category}：保留横杠后四位", category

    match = re.match(r"^(.+-[A-Za-z]+)\d+$", sku)
    if match:
        return match.group(1), "按鞋码格式兜底处理", category or "未识别"

    left, sep, right = sku.partition("-")
    if sep and len(right) >= 4:
        return f"{left}-{right[:4]}", "按横杠后四位兜底处理", category or "未识别"

    return sku, "未处理：无法识别格式", category or "未识别"


def decode_category(value: object) -> str:
    text = str(value or "").strip()
    known = {
        "Ьѥ": "鞋靴",
        "��װ": "服装",
        "���": "箱包",
    }
    if text in known:
        return known[text]
    try:
        decoded = text.encode("utf-8").decode("gbk")
        return known.get(decoded, decoded)
    except UnicodeError:
        return known.get(text, text)


def read_shop_rows() -> tuple[list[dict[str, str]], Counter[str]]:
    workbook = openpyxl.load_workbook(SHOP_FILE, read_only=True, data_only=True)
    sheet = workbook.active
    sheet.reset_dimensions()

    rows = sheet.iter_rows(values_only=True)
    next(rows, None)

    seen: set[str] = set()
    result: list[dict[str, str]] = []
    categories: Counter[str] = Counter()

    for excel_row, row in enumerate(rows, start=2):
        category = decode_category(row[2] if len(row) > 2 else "")
        sku = str(row[9] or "").strip() if len(row) > 9 else ""
        product_id = str(row[0] or "").strip() if len(row) > 0 else ""
        product_name = str(row[1] or "").strip() if len(row) > 1 else ""
        link = str(row[19] or "").strip() if len(row) > 19 else ""
        style_no = str(row[20] or "").strip() if len(row) > 20 else ""

        if category:
            categories[category] += 1
        if not sku:
            continue

        normalized, rule, fixed_category = normalize_shop_sku(sku, row[2] if len(row) > 2 else "")
        key = normalized or sku
        if key in seen:
            continue
        seen.add(key)

        result.append(
            {
                "excel_row": str(excel_row),
                "product_id": product_id,
                "category": fixed_category,
                "sku": sku,
                "normalized": normalized,
                "rule": rule,
                "style_no": style_no,
                "product_name": product_name,
                "link": link,
            }
        )

    return result, categories


def read_brand_codes() -> set[str]:
    workbook = openpyxl.load_workbook(BRAND_FILE, read_only=True, data_only=True)
    sheet = workbook.active
    codes: set[str] = set()
    for row in sheet.iter_rows(min_row=3, values_only=True):
        code = str(row[2] or "").strip() if len(row) > 2 else ""
        if code:
            codes.add(code)
    return codes


def render_html(rows: list[dict[str, str]], categories: Counter[str], brand_codes: set[str]) -> str:
    matched = [row for row in rows if row["normalized"] in brand_codes]
    unmatched = [row for row in rows if row["normalized"] not in brand_codes]
    preview_rows = rows[:500]

    body = "\n".join(render_row(row, brand_codes) for row in preview_rows)
    category_text = " / ".join(f"{escape(key)} {count}" for key, count in categories.most_common())

    return f"""<!doctype html>
<meta charset="utf-8">
<title>店铺 SKU 处理预览</title>
<style>
  :root {{
    font-family: "Microsoft YaHei", "Segoe UI", Arial, sans-serif;
    color: #1f2d3d;
    background: #eef3f8;
  }}
  body {{
    margin: 0;
    padding: 22px;
  }}
  .wrap {{
    max-width: 1360px;
    margin: 0 auto;
    background: #fff;
    border: 1px solid #d9e2ec;
    border-radius: 12px;
    overflow: hidden;
  }}
  header {{
    display: flex;
    justify-content: space-between;
    gap: 18px;
    padding: 18px 20px;
    border-bottom: 1px solid #d9e2ec;
  }}
  h1 {{
    margin: 0 0 6px;
    font-size: 20px;
  }}
  p {{
    margin: 0;
    color: #607086;
    font-size: 13px;
  }}
  .stats {{
    display: flex;
    gap: 10px;
    flex-wrap: wrap;
    justify-content: flex-end;
  }}
  .stat {{
    min-width: 120px;
    padding: 10px 12px;
    border: 1px solid #e0e8f0;
    border-radius: 8px;
    background: #f8fbff;
  }}
  .stat span {{
    display: block;
    color: #6b7a90;
    font-size: 12px;
  }}
  .stat strong {{
    display: block;
    margin-top: 4px;
    font-size: 20px;
  }}
  .note {{
    padding: 10px 20px;
    border-bottom: 1px solid #e4ebf2;
    color: #40536b;
    font-size: 13px;
  }}
  table {{
    width: 100%;
    border-collapse: collapse;
    font-size: 13px;
  }}
  th {{
    position: sticky;
    top: 0;
    z-index: 1;
    background: #f5f8fb;
    color: #526278;
    text-align: left;
    font-weight: 600;
  }}
  th, td {{
    padding: 10px 12px;
    border-bottom: 1px solid #e8eef5;
    vertical-align: top;
  }}
  tbody tr:hover {{
    background: #f8fbff;
  }}
  .code {{
    font-family: Consolas, "Courier New", monospace;
    white-space: nowrap;
  }}
  .ok {{
    color: #137a4f;
    font-weight: 600;
  }}
  .miss {{
    color: #a36405;
    font-weight: 600;
  }}
  .muted {{
    color: #7a889b;
  }}
  .name {{
    max-width: 360px;
  }}
</style>
<div class="wrap">
  <header>
    <div>
      <h1>店铺在售表 SKU 处理预览</h1>
      <p>原始表不修改；这里只演示分类、款号截取和品牌表匹配结果。</p>
    </div>
    <div class="stats">
      <div class="stat"><span>去重后款色</span><strong>{len(rows)}</strong></div>
      <div class="stat"><span>品牌表匹配</span><strong>{len(matched)}</strong></div>
      <div class="stat"><span>未匹配</span><strong>{len(unmatched)}</strong></div>
    </div>
  </header>
  <div class="note">分类统计：{category_text}。当前显示前 {len(preview_rows)} 个去重后的款色。</div>
  <table>
    <thead>
      <tr>
        <th>Excel 行</th>
        <th>分类</th>
        <th>原始商家 SKU 编码</th>
        <th>处理后匹配编码</th>
        <th>规则</th>
        <th>品牌总表</th>
        <th>货号</th>
        <th>商品名</th>
      </tr>
    </thead>
    <tbody>{body}</tbody>
  </table>
</div>
"""


def render_row(row: dict[str, str], brand_codes: set[str]) -> str:
    status = "匹配" if row["normalized"] in brand_codes else "未匹配"
    status_class = "ok" if status == "匹配" else "miss"
    return f"""
      <tr>
        <td>{escape(row["excel_row"])}</td>
        <td>{escape(row["category"])}</td>
        <td class="code">{escape(row["sku"])}</td>
        <td class="code">{escape(row["normalized"])}</td>
        <td>{escape(row["rule"])}</td>
        <td class="{status_class}">{status}</td>
        <td class="code muted">{escape(row["style_no"])}</td>
        <td class="name">{escape(row["product_name"])}</td>
      </tr>"""


def escape(value: object) -> str:
    return html.escape(str(value or ""))


def self_check() -> None:
    assert normalize_shop_sku("246240-YEL12", "Ьѥ")[0] == "246240-YEL"
    assert normalize_shop_sku("L326U048-001899", "��װ")[0] == "L326U048-0018"
    assert normalize_shop_sku("L326U048-001899", "���")[0] == "L326U048-0018"


def main() -> None:
    self_check()
    rows, categories = read_shop_rows()
    brand_codes = read_brand_codes()
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_FILE.write_text(render_html(rows, categories, brand_codes), encoding="utf-8")
    matched = sum(1 for row in rows if row["normalized"] in brand_codes)
    print(f"rows={len(rows)} matched={matched} output={OUTPUT_FILE}")


if __name__ == "__main__":
    main()
